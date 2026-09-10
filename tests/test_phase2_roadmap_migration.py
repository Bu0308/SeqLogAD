"""P2.0.1 authority, provenance and fail-closed execution guards."""

import hashlib
import json
from pathlib import Path

import openpyxl
import yaml

ROOT = Path(__file__).resolve().parents[1]
ORDER = ["P2.PRE"] + [f"P2.{i}" for i in range(1, 8)]


def read(path):
    return yaml.safe_load((ROOT / path).read_text())


def sha(path):
    return hashlib.sha256((ROOT / path).read_bytes()).hexdigest()


def cells(path):
    workbook = openpyxl.load_workbook(ROOT / path, data_only=False)
    return {f"{s.title}!{c.coordinate}": c.value
            for s in workbook for row in s for c in row if c.value is not None}


P2 = read("configs/protocols/phase2-architecture-v1.yaml")["phase2"]
MIG = read("configs/protocols/phase2-roadmap-migration-v1.yaml")["migration"]
BASE = read("configs/models/base-freeze-v1.yaml")["base_freeze"]
MIGRATED_WORKBOOK = BASE["roadmap_status_update"]["previous_workbook"]


def test_unique_ids_and_exact_linear_order_across_authorities():
    node = yaml.compose((ROOT / "configs/protocols/phase2-architecture-v1.yaml").read_text())
    phase2_node = next(value for key, value in node.value if key.value == "phase2")
    tasks_node = next(value for key, value in phase2_node.value if key.value == "tasks")
    raw_ids = [key.value for key, _ in tasks_node.value]
    assert len(raw_ids) == len(set(raw_ids))  # safe_load would hide duplicate mapping keys.
    roadmap = read("configs/plan/excel-roadmap-v1.yaml")["roadmap"]
    ids = [t["task_id"] for t in roadmap["tasks"]]
    assert len(ids) == len(set(ids)) == 32
    tasks = [t for t in roadmap["tasks"] if t["task_id"].startswith("P2.")]
    assert [t["task_id"] for t in tasks] == list(P2["tasks"]) == MIG["execution_order"] == ORDER
    for index, task in enumerate(tasks):
        expected = ["P1.8"] if index == 0 else [ORDER[index - 1]]
        contract = P2["tasks"][task["task_id"]]
        assert task["dependencies"] == contract["depends_on"] == expected
        assert task["workstream"] == contract["name"]
        assert task["workbook_status"] == ("Done" if index == 0 else "Not started")
    for task in roadmap["tasks"]:
        assert set(task["dependencies"]) <= set(ids)
    for path in (".ai/task-routing.md", ".ai/primary-context.md", "README.md"):
        assert "NEXT_AUTHORIZED_TASK = P2.1" in (ROOT / path).read_text()
    routing = (ROOT / ".ai/task-routing.md").read_text()
    assert " → ".join(ORDER) in routing


def test_every_old_phase2_id_has_an_unambiguous_historical_crosswalk():
    assert MIG["crosswalk"] == {
        "P2.1": ["P2.PRE"], "P2.2": ["P2.1"], "P2.3": ["P2.1"],
        "P2.4": ["P2.1"], "P2.5": ["P2.2"], "P2.6": ["P2.2"],
        "P2.7": ["P2.4"], "P2.8": ["P2.5"],
    }
    old = read(MIG["previous_projection"])["roadmap"]["tasks"]
    assert {t["task_id"] for t in old if t["task_id"].startswith("P2.")} == set(MIG["crosswalk"])
    for old_id, new_ids in MIG["crosswalk"].items():
        assert all(old_id in P2["tasks"][tid]["old_tasks"] for tid in new_ids)
    assert "conditional" in MIG["crosswalk_notes"]["P2.3"]
    assert "P2.6" in MIG["crosswalk_notes"]["P3.1"]


def test_migration_preserves_signed_history_and_chains_exact_hashes():
    previous = read("configs/protocols/g0-signatures.yaml")["g0_signatures"]["plan_amendments"][0]
    assert previous["to_sha256"] == MIG["from_sha256"] == sha(MIG["previous_workbook"])
    assert MIG["to_sha256"] == sha(MIGRATED_WORKBOOK)
    assert MIG["previous_projection_sha256"] == sha(MIG["previous_projection"])
    for path, digest in MIG["frozen_history"].items():
        assert sha(path) == digest
    assert MIG["phase1_contracts_changed"] is MIG["gate_signature_created"] is False
    receipt = json.loads((ROOT / "data/processed/protocol/G0-receipt.json").read_text())
    for old_path, archive in MIG["historical_receipt_bindings"].items():
        assert sha(archive) == receipt["artifact_sha256"][old_path]


def test_only_declared_phase2_and_dependent_cells_changed():
    before, after = cells(MIG["previous_workbook"]), cells(MIGRATED_WORKBOOK)
    moved = {k for k in before.keys() | after.keys() if before.get(k) != after.get(k)}
    assert moved == set(MIG["changed_cells"])
    permitted = {f"Task Register!{c}{r}" for r in range(13, 21) for c in "ABCDEFGHIJK"}
    permitted |= {f"Task Playbooks!{c}{r}" for r in range(13, 21) for c in "ABCDE"}
    permitted |= {"Task Register!E21", "Task Register!F21", "Task Register!I21",
                  "Task Register!J21", "Task Register!I22", "Task Register!I25",
                  "Task Playbooks!B21", "Task Playbooks!C21", "Task Playbooks!D21",
                  "Gates!F7", "Executive!B11", "Executive!E11", "Roadmap!C6", "Roadmap!E6"}
    assert moved <= permitted  # No P1, G0, research-source or unrelated P3/P4 edits.
    old = openpyxl.load_workbook(ROOT / MIG["previous_workbook"])
    new = openpyxl.load_workbook(ROOT / MIGRATED_WORKBOOK)
    assert old.sheetnames == new.sheetnames
    for s in old:
        assert str(s.merged_cells) == str(new[s.title].merged_cells)
        assert set(s.tables) == set(new[s.title].tables)
        assert str(s.data_validations) == str(new[s.title].data_validations)


def test_data_notebooks_experts_and_label_permissions_are_unchanged():
    for key, digest in MIG["preserved_p20_section_sha256"].items():
        assert hashlib.sha256(json.dumps(P2[key], sort_keys=True).encode()).hexdigest() == digest
    assert P2["portability"]["data_root"] == "data/phase2"
    assert P2["portability"]["fold_root"] == "folds/{fold_id}"
    assert [e["notebook"] for e in P2["experts"]] == [
        "notebooks/P2.1_train_semantic_expert.ipynb",
        "notebooks/P2.2_train_sequence_reference.ipynb",
        "notebooks/P2.3_train_lightweight_sequence_expert.ipynb",
        "notebooks/P2.4_train_gtat_expert.ipynb",
    ]
    assert P2["source_label_scope"] == "UNRESOLVED_DENY"
    assert P2["source_label_blocker"]["training_or_selection_permission"] is False
    for tid in ("P2.6", "P2.7"):
        assert "SOURCE_LABEL_SCOPE_FOR_SUPERVISED_METRICS" in P2["tasks"][tid]["blockers"]


def test_base_prerequisite_pass_is_metadata_only():
    base = read(P2["base_freeze_contract"])["base_freeze"]
    assert base["task_id"] == "P2.PRE"
    assert base["status"] == "PASS"
    assert base["completion_status"] == "METADATA_VERIFIED"
    assert base["execution_scope"] == "METADATA_ONLY"
    assert base["download_authorized"] is base["training_authorized"] is False
    fields = base["required_metadata"]
    assert {"base_model_identifier", "model_revision_commit", "tokenizer_identifier",
            "tokenizer_revision_commit", "quantization_policy", "precision_policy",
            "peft_lora_library_constraints", "expected_context_length",
            "license_identifier_and_terms_url", "access_requirements_and_verification",
            "artifact_checkpoint_naming", "reproducibility_metadata"} <= set(fields)
    assert fields["base_model_identifier"] == "meta-llama/Llama-3.1-8B"
    assert fields["model_revision_commit"] == "d04e592bb4f6aa9cfee91e2e20afa771667e1d4b"
    assert fields["tokenizer_identifier"] == "meta-llama/Llama-3.1-8B"
    assert fields["tokenizer_revision_commit"] == "d04e592bb4f6aa9cfee91e2e20afa771667e1d4b"

    policy_refs = {
        "quantization_policy": "#quantization",
        "precision_policy": "#precision",
        "peft_lora_library_constraints": "#runtime",
        "expected_context_length": "#context",
        "license_identifier_and_terms_url": "#access",
        "access_requirements_and_verification": "#access",
        "artifact_checkpoint_naming": "#artifacts",
        "reproducibility_metadata": "#artifacts",
        "gpu_memory_and_fair_comparison_budget": "#hardware",
    }

    for key, expected in policy_refs.items():
        assert fields[key] == expected
    assert base["review"]["status"] == "PASS_METADATA_VERIFIED_RUNTIME_PREFLIGHT_DEFERRED"
    assert base["future_runtime_preflight"]["owner"] == "P2.1"
    assert base["future_runtime_preflight"]["executed"] is False
    state = read("configs/active-state.yaml")["active_state"]
    assert state["next_authorized_task"] == "P2.1"
    assert state["semantic_execution_authorized"] is True
    assert state["phase3_enabled"] is False
    assert P2["tasks"]["P2.1"]["execution_authorized"] is True
    assert P2["training_authorized"] is P2["phase3_enabled"] is False
    assert all(state["gates"][g] == "NOT_PASSED" for g in ("G1", "G2", "G3", "G4"))


def test_future_receipt_routes_metadata_prerequisite_without_rewriting_history():
    from seqlogad.protocol.freeze import build_g0_receipt

    path = "data/processed/protocol/G0-receipt.json"
    before = sha(path)
    receipt = build_g0_receipt(ROOT)  # Read-only derivation; no build/write command.
    historical = json.loads((ROOT / path).read_text())
    assert receipt["criteria"] == historical["criteria"]  # All 19 P1 criteria unchanged.
    assert receipt["gate_state"] == "PROTOCOL_READY"
    assert receipt["next_authorized_task"] == "P2.PRE"
    assert receipt["unlocked_by_this_gate"] == [
        "P2.PRE (metadata freeze only; expert execution requires its verified completion)"
    ]
    assert set(receipt["still_gated"]) == {"G1", "G2", "G3", "G4"}
    assert sha(path) == before == MIG["frozen_history"][path]
