"""Metadata freeze guards; these tests do not load models or validate GPU support."""
import hashlib
import json
import re
from pathlib import Path

import openpyxl
import yaml

ROOT = Path(__file__).resolve().parents[1]


def read(path):
    return yaml.safe_load((ROOT / path).read_text())


BASE = read('configs/models/base-freeze-v1.yaml')['base_freeze']
P2 = read('configs/protocols/phase2-architecture-v1.yaml')['phase2']
STATE = read('configs/active-state.yaml')['active_state']


def test_immutable_identity_is_bound_to_recorded_upstream_evidence():
    raw = (ROOT / BASE['evidence']['path']).read_bytes()
    assert hashlib.sha256(raw).hexdigest() == BASE['evidence']['sha256']
    observed = json.loads(raw)['model_api']['observed']
    fields = BASE['required_metadata']
    assert fields['base_model_identifier'] == fields['tokenizer_identifier'] == observed['id']
    assert fields['model_revision_commit'] == fields['tokenizer_revision_commit'] == observed['sha']
    assert re.fullmatch('[a-f0-9]{40}', observed['sha'])
    assert BASE['revision_policy']['allow_floating_revision'] is False
    assert (ROOT / BASE['document']).is_file()


def test_authenticated_metadata_is_required_for_pass_and_only_semantic_is_authorized():
    assert BASE['status'] == P2['tasks']['P2.PRE']['status'] == 'PASS'
    assert BASE['blockers'] == BASE['unresolved_fields'] == []
    assert BASE['access']['owner_confirmation'] == 'USER_CONFIRMED_HF_ACCESS_GRANTED'
    assert BASE['access']['authenticated_metadata_verified'] is True
    verification = BASE['authenticated_verification']
    assert verification['gated_access'] == 'PASS'
    revision = BASE['required_metadata']['model_revision_commit']
    assert verification['revision'] == verification['revision_api']['observed_sha'] == revision
    configs = {}
    for name, item in verification['files'].items():
        raw = (ROOT / item['path']).read_bytes()
        assert hashlib.sha256(raw).hexdigest() == item['sha256']
        assert item['revision'] == revision and f'/resolve/{revision}/' in item['url']
        assert item['http_status'] == 200
        configs[name] = json.loads(raw)
    assert set(configs) == {'config.json', 'tokenizer_config.json'}
    assert BASE['revision_policy']['model_config_sha256'] == verification['files']['config.json']['sha256']
    assert BASE['revision_policy']['tokenizer_config_sha256'] == verification['files']['tokenizer_config.json']['sha256']
    c, t = configs['config.json'], configs['tokenizer_config.json']
    assert t['tokenizer_class'] == BASE['tokenizer']['resolved_class'] == 'PreTrainedTokenizerFast'
    assert 'auto_map' not in t and 'pad_token' not in t
    assert c['max_position_embeddings'] == t['model_max_length'] == BASE['context']['upstream_capability_tokens']
    for role in ('bos', 'eos'):
        token_id = c[f'{role}_token_id']
        assert token_id == BASE['tokenizer']['numeric_special_token_ids'][role]
        assert t['added_tokens_decoder'][str(token_id)]['content'] == t[f'{role}_token']
    assert verification['model_weights_downloaded'] is False
    assert P2['tasks']['P2.1']['depends_on'] == ['P2.PRE']
    assert P2['tasks']['P2.1']['execution_authorized'] is STATE['semantic_execution_authorized'] is True
    assert P2['tasks']['P2.1']['status'] == 'NOT_STARTED'
    assert all(not P2['tasks'][f'P2.{i}']['execution_authorized'] for i in range(2, 8))
    assert STATE['phase3_enabled'] is P2['phase3_enabled'] is False
    assert BASE['data']['source_label_scope'] == P2['source_label_scope'] == 'UNRESOLVED_DENY'


def test_base_is_shared_but_adapter_and_optimizer_states_are_isolated():
    assert BASE['model']['base_weights_immutable'] is True
    assert BASE['model']['same_upstream_base_for_both_adapters'] is True
    assert BASE['peft']['adapter_isolation'] is True
    for key in ('share_adapter_weights', 'share_optimizer_or_resume_state',
                'trainable_base_parameters', 'trainable_embeddings'):
        assert BASE['peft'][key] is False
    assert BASE['peft']['bias'] == 'none'
    assert all(v is None for v in BASE['peft']['experimental_hyperparameters'].values())
    assert BASE['tokenizer']['new_eventid_tokens_allowed'] is False
    assert BASE['tokenizer']['resize_embeddings_allowed'] is False


def test_precision_and_quantization_are_explicit_without_measurement_claims():
    assert 'including_emulation=False' in BASE['precision']['native_check']
    assert BASE['precision']['fallback'] == 'FP16_CONDITIONAL_ON_FINITE_PREFLIGHT'
    assert BASE['quantization']['default_training_mode'] == 'QLORA_NF4_DOUBLE_QUANT'
    assert BASE['quantization']['low_memory_mode'] == 'QLORA_NF4_DOUBLE_QUANT_LOW_MEMORY'
    assert BASE['quantization']['bnb_4bit_quant_type'] == 'nf4'
    assert BASE['quantization']['bnb_4bit_use_double_quant'] is True
    assert BASE['hardware']['measurement_status'] == 'NOT_MEASURED_BY_SEQLOGAD'
    assert BASE['hardware']['minimum_is_guaranteed'] is False
    assert BASE['future_runtime_preflight']['executed'] is False
    assert BASE['context']['actual_horizon'] is None


def test_colab_contract_preserves_data_and_future_notebooks():
    notebook = BASE['notebook']
    assert notebook['primary_training_environment'] == 'Google Colab'
    assert notebook['colab_data_root'] == '/content/phase2'
    assert notebook['applies_to_every_phase2_training_task'] is True
    assert notebook['colab_ready_ipynb_required'] is True
    assert notebook['required_future_path'] == 'notebooks/P2.1_train_semantic_expert.ipynb'
    assert BASE['data']['root'] == P2['portability']['data_root'] == 'data/phase2'
    for view in ('semantic', 'sequence', 'graph'):
        assert BASE['data'][view] == f'data/phase2/folds/{{fold_id}}/{view}'
    assert not (ROOT / BASE['data']['root']).exists()
    assert all(not (ROOT / expert['notebook']).exists() for expert in P2['experts'])
    assert BASE['artifacts']['output_root'] == 'outputs/phase2'
    assert {'base_model_revision', 'tokenizer_revision', 'data_manifest_sha256',
            'seqlogad_git_commit', 'hardware'} <= set(BASE['artifacts']['manifest_fields'])


def test_workbook_status_update_has_exact_provenance_and_no_other_cell_changes():
    update = BASE['roadmap_status_update']
    migration = read('configs/protocols/phase2-roadmap-migration-v1.yaml')['migration']
    before_path = ROOT / update['previous_workbook']
    after_path = ROOT / update['result_workbook']
    assert hashlib.sha256(before_path.read_bytes()).hexdigest() == update['from_sha256'] == migration['to_sha256']
    assert hashlib.sha256(after_path.read_bytes()).hexdigest() == update['to_sha256']
    workbooks = [openpyxl.load_workbook(p, data_only=False) for p in (before_path, after_path)]
    before, after = [{f'{s.title}!{c.coordinate}': c.value for s in w for row in s for c in row
                      if c.value is not None} for w in workbooks]
    changed = {k: {'before': before.get(k), 'after': after.get(k)}
               for k in before.keys() | after.keys() if before.get(k) != after.get(k)}
    assert changed == update['changed_cells'] == {'Task Register!K13': {'before': 'Not started', 'after': 'Blocked'}}
    assert workbooks[0].sheetnames == workbooks[1].sheetnames
    for sheet in workbooks[0]:
        current = workbooks[1][sheet.title]
        assert str(sheet.merged_cells) == str(current.merged_cells)
        assert set(sheet.tables) == set(current.tables)
        assert str(sheet.data_validations) == str(current.data_validations)
    assert update['phase1_changed'] is update['gate_signature_created'] is False


def test_runtime_pins_match_verified_direct_dependency_metadata():
    from packaging.requirements import Requirement
    from packaging.specifiers import SpecifierSet
    from packaging.utils import canonicalize_name

    evidence = json.loads((ROOT / BASE['evidence']['path']).read_text())
    pins = {canonicalize_name(k): v for k, v in BASE['runtime']['pins'].items()}
    packages = {canonicalize_name(p['name']): p for p in evidence['packages']}
    for name, version in pins.items():
        package = packages[name]
        assert package['version'] == version
        assert '3.12' in SpecifierSet(package['requires_python'])
        for text in package['active_requires_dist_linux_py312']:
            requirement = Requirement(text)
            dependency = canonicalize_name(requirement.name)
            if dependency in pins:
                assert pins[dependency] in requirement.specifier
    # This checks recorded direct constraints, not an installed/transitive GPU stack.
    assert BASE['runtime']['local_environment_modified'] is False


def test_completion_status_update_preserves_blocked_history():
    update = BASE['completion_status_update']
    before_path = ROOT / update['previous_workbook']
    after_path = ROOT / 'Bang_ke_hoach_SeqLogAD.xlsx'
    assert hashlib.sha256(before_path.read_bytes()).hexdigest() == update['from_sha256'] == BASE['roadmap_status_update']['to_sha256']
    assert hashlib.sha256(after_path.read_bytes()).hexdigest() == update['to_sha256'] == STATE['authoritative_plan']['sha256']
    snapshots = [{f'{s.title}!{c.coordinate}': c.value for s in openpyxl.load_workbook(p)
                  for row in s for c in row if c.value is not None} for p in (before_path, after_path)]
    before, after = snapshots
    changed = {k: {'before': before.get(k), 'after': after.get(k)}
               for k in before.keys() | after.keys() if before.get(k) != after.get(k)}
    assert changed == update['changed_cells'] == {'Task Register!K13': {'before': 'Blocked', 'after': 'Done'}}
    assert update['phase1_changed'] is update['gate_signature_created'] is False
